from __future__ import annotations

import csv
import sys
import types
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = REPO_ROOT.parents[1] / "outputs"
DOCX_PATH = OUTPUT_DIR / "mumbai_public_transport_accessibility_research_paper.docx"


def _add_bundled_site_packages() -> None:
    if str(REPO_ROOT) not in sys.path:
        sys.path.insert(0, str(REPO_ROOT))
    candidate = Path(sys.executable).parent / "Lib" / "site-packages"
    if candidate.is_dir() and str(candidate) not in sys.path:
        sys.path.insert(0, str(candidate))


def _install_import_stubs_if_needed() -> None:
    try:
        import flask  # noqa: F401
    except ModuleNotFoundError:
        flask = types.ModuleType("flask")
        flask.Flask = type(
            "Flask",
            (),
            {
                "__init__": lambda self, *args, **kwargs: None,
                "route": lambda self, *args, **kwargs: (lambda fn: fn),
                "run": lambda self, *args, **kwargs: None,
            },
        )
        flask.render_template = lambda *args, **kwargs: kwargs
        flask.request = types.SimpleNamespace(
            get_json=lambda *args, **kwargs: {}
        )
        flask.jsonify = lambda value: value
        sys.modules["flask"] = flask

    try:
        import scipy  # noqa: F401
    except ModuleNotFoundError:
        scipy = types.ModuleType("scipy")
        spatial = types.ModuleType("scipy.spatial")
        spatial.cKDTree = object
        scipy.spatial = spatial
        sys.modules["scipy"] = scipy
        sys.modules["scipy.spatial"] = spatial


def load_project_results():
    _add_bundled_site_packages()
    _install_import_stubs_if_needed()
    import pandas as pd
    import app

    data_dir = REPO_ROOT / "data"
    dataset_files = {
        "BMC ward population": (
            "bmc_ward_population.csv",
            "Population input for area-priority screening",
        ),
        "Railway train services": (
            "mumbai_local_train_services_ALL_OFFICIAL.csv",
            "Official uploaded Indian Railways timetable service records",
        ),
        "Railway station-line records": (
            "mumbai_local_train_stations_ALL_OFFICIAL.csv",
            "Station and line structure for the timetable graph",
        ),
        "Railway stop-time records": (
            "mumbai_local_train_stop_times_ALL_OFFICIAL.csv",
            "Official timetable stop sequences and scheduled times",
        ),
        "Verified BEST routes": (
            "best_official_verified_routes.csv",
            "BEST/BMC route references with official MCGM/BMC sources",
        ),
        "BEST summary metrics": (
            "best_official_network_summary.csv",
            "BEST/BMC network statistics with official MCGM source URLs",
        ),
    }
    dataset_rows = []
    for label, (filename, role) in dataset_files.items():
        frame = pd.read_csv(data_dir / filename, encoding="utf-8-sig")
        dataset_rows.append((label, len(frame), f"{filename} — {role}"))

    stop_times = pd.read_csv(
        data_dir / "mumbai_local_train_stop_times_ALL_OFFICIAL.csv",
        encoding="utf-8-sig",
    )
    service_sequences = len(
        stop_times[
            ["line", "train_id", "source_file", "source_page"]
        ].drop_duplicates()
    )

    return app, dataset_rows, service_sequences


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", fill)


def _set_cell_text(cell, text: str, bold: bool = False, color: str = "171717") -> None:
    from docx.shared import Pt, RGBColor

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = 0
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(color)


def _add_table(document, headers, rows):
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    for cell, header in zip(header_cells, headers):
        _set_cell_shading(cell, "171717")
        _set_cell_text(cell, header, bold=True, color="FFFFFF")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            _set_cell_text(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()
    return table


def _add_bullet(document, text: str) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def _add_body(document, text: str) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08


def _add_heading(document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def create_paper() -> Path:
    _add_bundled_site_packages()
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    app, dataset_rows, service_sequences = load_project_results()
    results = app.cached_algorithm_results
    graph_nodes = len(app.cached_graph)
    graph_edges = sum(len(edges) for edges in app.cached_graph.values())
    benchmark_routes = results[0]["benchmark_routes"] if results else 0
    leader = results[0] if results else {}
    runner_up = results[1] if len(results) > 1 else {}

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.color.rgb = RGBColor(5, 150, 105)
    styles["Heading 2"].font.name = "Aptos Display"

    document.core_properties.title = "Mumbai Public Transport Accessibility"
    document.core_properties.author = "Krishna Tiwari; Yash Karande; Sushant Eppakayala; Vedant Nayak"
    document.core_properties.subject = "PBL study of public-transport priority areas and timetable graph evaluation"
    document.core_properties.keywords = "Mumbai, public transport, transit deserts, priority areas, graph search, accessibility"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Mumbai Public Transport Accessibility: \nPriority Areas, Transit-Desert Indicators, and Route Evaluation")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run("Population-based prioritization and timetable-graph evaluation for better connectivity")
    run.italic = True
    run.font.color.rgb = RGBColor(82, 82, 82)

    byline = document.add_paragraph()
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.add_run(
        "Krishna Tiwari · Yash Karande · Sushant Eppakayala · Vedant Nayak\n23 September 2026"
    ).bold = True
    document.add_paragraph()

    _add_heading(document, "Abstract", 1)
    _add_body(
        document,
        f"This PBL project addresses Mumbai public-transport accessibility by using population and public-transport network data to identify areas that could be prioritized for better connectivity. The system combines a ward-level population-priority indicator with official railway timetable-derived graph data and BEST/BMC route references. Five graph-search algorithms—Breadth-First Search, Depth-First Search, Uniform Cost Search, Greedy Best-First Search, and A* Search—are evaluated on 300 deterministic, reachable station pairs. The study reports success rate, route optimality, search efficiency, compute-time efficiency, path cost, expanded nodes, a transparent evaluation rank, and regression-style route-cost errors against the best measured route. {leader.get('algorithm', 'The leading algorithm')} ranked first under the lexicographic rule Winner = arg min (route-cost RMSE, nodes checked, compute time). Because the supplied ward file has population but no ward coordinates or transit-distance field, the area result is a screening indicator for possible transit-priority areas, not proof of a geographic transit desert. The regression metrics are measurement-based diagnostics, not evidence of a trained machine-learning model."
    )
    _add_body(document, "Keywords: Mumbai public transport; priority areas; transit deserts; graph search; route optimization; timetable graph; accessibility planning")

    _add_heading(document, "1. Introduction", 1)
    _add_body(
        document,
        "The PBL problem is to use public-transport station, route, population, and location information to identify areas with poor access to public transportation and suggest areas that could be prioritized for better connectivity. Urban public-transport systems are naturally represented as networks of stations, services, and timed connections, so the project combines an area-priority view with an auditable timetable-graph evaluation."
    )
    _add_body(
        document,
        "The research questions are: which population areas should be screened first for improved connectivity, and how do alternative graph-search algorithms differ when evaluated on the same official timetable network and reachable station pairs? The contribution is a transparent PBL workflow that connects area prioritization with route-search evidence while making algorithm metrics visible."
    )

    _add_heading(document, "2. Data and Network Construction", 1)
    _add_body(
        document,
        "The application loads official CSV files committed with the repository. The dataset collection includes the BMC ward-population file; the railway timetable dataset (train services, station-line records, and stop-time records); verified BEST routes; and BEST network summary metrics. The railway stop-time records identify Indian Railways and official uploaded railway timetable PDFs in their source fields. The BEST route and summary files cite official MCGM/BMC documents. The railway CSVs do not contain a separate MCGM source URL, so this paper describes them as official timetable-derived railway data rather than attributing them to MCGM without evidence."
    )
    _add_body(
        document,
        "For graph construction, station stop sequences are grouped by line and train identifier, sorted by stop sequence, and converted into directed edges. Each edge weight is the scheduled time difference between consecutive stops. Overnight crossings are adjusted by one day, and non-positive intervals are excluded rather than invented."
    )
    _add_table(
        document,
        ["Dataset", "Rows", "Role"],
        [(label, count, filename) for label, count, filename in dataset_rows],
    )
    _add_body(
        document,
        f"The processed graph contains {graph_nodes} station nodes and {graph_edges} directed timetable connections derived from {service_sequences:,} official timetable service sequences. The benchmark samples {benchmark_routes} reachable origin-destination pairs with a fixed random seed of 42, ensuring that every algorithm is compared on the same route set."
    )

    _add_heading(document, "3. Algorithms", 1)
    _add_body(document, "The implementation evaluates the following algorithms:")
    for description in (
        "BFS explores the graph layer by layer and prioritizes hop depth rather than scheduled travel time.",
        "DFS follows one branch deeply before backtracking and does not guarantee a least-cost route.",
        "Uniform Cost Search expands the lowest accumulated scheduled cost and provides the cost reference for non-negative edge weights.",
        "Greedy Best-First Search prioritizes a topology-derived estimate of remaining hops and may sacrifice route quality for search speed.",
        "A* Search combines accumulated cost with the same admissible topology-derived heuristic."
    ):
        _add_bullet(document, description)

    _add_heading(document, "4. Evaluation Metrics", 1)
    _add_body(document, "The benchmark reports normalized metrics where higher values are better. For an algorithm a and a benchmark route r:")
    _add_body(document, "Route optimality = best observed route cost / route cost produced by a")
    _add_body(document, "Search efficiency = minimum expanded nodes / expanded nodes used by a")
    _add_body(document, "Compute Time = measured algorithm execution duration in seconds")
    _add_body(document, "Compute-Time efficiency = fastest successful compute time / compute time of a")
    _add_body(document, "Evaluation ranking: Winner = arg min (route-cost RMSE, nodes checked, compute time). This means the lowest route-cost RMSE wins; fewer nodes checked wins when RMSE is tied, and compute time is considered only when both earlier values are tied.")
    _add_body(
        document,
        "Success rate is the number of successful searches divided by the number of valid benchmark routes. Search efficiency and compute-time efficiency remain visible supporting measurements, while the evaluation rank uses the explicit lexicographic rule rather than an opaque weighted composite score."
    )
    _add_body(
        document,
        "Because route cost is a continuous measured quantity, the study also reports regression-style diagnostics. For each algorithm, its measured route cost is compared with the best measured route cost for the same origin-destination pair. MAE is the mean absolute error in minutes, MSE is the mean squared error in minutes squared, RMSE is the root mean squared error in minutes, MAPE is the mean absolute percentage error, and R-squared summarizes agreement with the reference costs."
    )

    _add_heading(document, "5. Results", 1)
    result_rows = []
    for result in results:
        result_rows.append(
            (
                result["algorithm"],
                f"{result['success_rate']:.1f}%",
                f"{result['optimality']:.1f}%",
                f"{result['efficiency']:.1f}%",
                f"{result['runtime_efficiency']:.1f}%",
                str(result["evaluation_rank"]),
                f"{result['path_cost']:.3f}",
                f"{result['nodes']:.2f}",
                f"{result['time']:.6f}",
            )
        )
    _add_table(
        document,
        [
            "Algorithm", "Success Rate", "Route Optimality", "Search Efficiency",
            "Compute-Time Efficiency", "Evaluation Rank", "Avg Cost (min)",
            "Nodes Checked", "Median Time (s)"
        ],
        result_rows,
    )
    _add_body(
        document,
        f"{leader.get('algorithm', 'The leading algorithm')} ranked first (evaluation rank {leader.get('evaluation_rank', 'n/a')}) with route-cost RMSE {leader.get('route_rmse', 0.0):.3f} minutes. The next-ranked method was {runner_up.get('algorithm', 'not available')}; tied methods were compared lexicographically by route-cost RMSE, nodes checked, and compute time. Greedy Best-First Search expanded the fewest nodes but produced routes with RMSE {next((r['route_rmse'] for r in results if r['algorithm'] == 'Greedy Best First Search'), 0.0):.3f} minutes. BFS had the same route-cost error profile as Greedy Best-First Search in this benchmark, while DFS produced the largest route-cost error."
    )

    _add_heading(document, "6. Regression-Style Route-Cost Evaluation", 1)
    regression_rows = []
    for result in results:
        regression_rows.append(
            (
                result["algorithm"],
                f"{result['route_mae']:.3f}",
                f"{result['route_mse']:.3f}",
                f"{result['route_rmse']:.3f}",
                f"{result['route_mape']:.1f}%",
                f"{result['route_r2']:.3f}",
                str(result["regression_samples"]),
            )
        )
    _add_table(
        document,
        ["Algorithm", "MAE (min)", "MSE (min²)", "RMSE (min)", "MAPE", "R-squared", "Samples"],
        regression_rows,
    )
    _add_body(
        document,
        "Uniform Cost Search and A* Search match the best measured route cost on this benchmark, so their route-cost errors are zero. Greedy Best-First Search and BFS find routes for every pair but can be substantially more expensive, which produces larger error values and negative R-squared values. These figures quantify route-cost agreement; they do not mean that the application trained a regression model."
    )

    _add_heading(document, "7. Priority-Area and Transit-Desert Indicator", 1)
    _add_body(
        document,
        "The PBL planning objective is to identify areas with poor access and suggest areas that could be prioritized for better connectivity. The repository therefore includes a ward-level planning view. Because the supplied ward file contains population but no ward coordinates or transit-distance field, the application does not claim to calculate a geographic transit-desert score. Instead, it reports a transparent population priority score equal to ward population divided by the maximum ward population, multiplied by 100. High-scoring wards are candidates for further access analysis and investment screening; they are not automatically proven transit deserts. A true geographic transit-desert study would require ward geometry, station coordinates, route distance or travel time, and service-frequency measures."
    )

    _add_heading(document, "8. Discussion", 1)
    _add_body(
        document,
        "The results show the central trade-off in route search. Uniform Cost Search and A* Search preserve the best observed scheduled travel cost, while Greedy Best-First Search and BFS reduce route-search effort or compute time at the cost of longer routes. DFS is sensitive to traversal order and can return a substantially more expensive path. The best algorithm therefore depends on the research objective: cost-optimal routing favors UCS or A*, while rapid exploratory search may favor Greedy Best-First Search."
    )
    _add_body(
        document,
        "Compute-time efficiency should be interpreted cautiously because the measured times are very small and can be affected by machine load. Compute time and search efficiency are therefore reported as supporting measurements, while the transparent evaluation rank uses compute time only after route-cost RMSE and nodes checked have been minimized."
    )

    _add_heading(document, "9. Limitations and Threats to Validity", 1)
    for limitation in (
        "The graph represents scheduled local-train stop sequences and does not model delays, crowding, disruptions, fares, platform changes, or walking transfers.",
        "The benchmark consists of 300 reachable station pairs and is deterministic but not a complete enumeration of all possible journeys.",
        "Compute time is hardware- and load-dependent, so comparisons should be repeated under controlled conditions for a formal performance study.",
        "The ward planning indicator screens population priority but cannot establish transit access because the provided ward data lacks geographic distance to stations.",
        "The project does not train a supervised regression model; the MAE, RMSE, MAPE, and R-squared values are derived route-cost diagnostics against the best measured route on each benchmark pair."
    ):
        _add_bullet(document, limitation)

    _add_heading(document, "10. Conclusion and Future Work", 1)
    _add_body(
        document,
        "This project provides a PBL-oriented screening workflow for Mumbai public-transport priority areas together with a transparent local-train graph-search benchmark. The leading algorithm ranked first under the explicit rule Winner = arg min (route-cost RMSE, nodes checked, compute time). The population-priority indicator helps identify where further connectivity analysis should begin, while the explicit metric panel demonstrates the route-search trade-offs in a reproducible research or classroom setting."
    )
    _add_body(
        document,
        "Future work should add live service disruption data, walking and interchange edges, controlled repeated compute-time experiments, ward geometries, geographic ward-to-station distances, service frequency, and a validated transit-desert index. These additions would allow the population screening indicator to become a true spatial accessibility analysis. A predictive or clustering study should be introduced only if a separate target dataset and research question are defined."
    )

    _add_heading(document, "References", 1)
    _add_body(document, "[1] ZeuS0nTopXD, Mumbai Public Transport Accessibility, GitHub repository, https://github.com/ZeuS0nTopXD/mumbai-public-transport-accessibility")
    _add_body(document, "[2] Official uploaded Indian Railways timetable PDF records listed in the repository railway service and stop-time CSV source fields.")
    source_urls = []
    with (REPO_ROOT / "data" / "best_official_network_summary.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            url = row.get("official_source", "").strip()
            if url and url not in source_urls:
                source_urls.append(url)
    for index, url in enumerate(source_urls[:1], start=3):
        _add_body(document, f"[{index}] Official BEST/BMC source referenced by the repository data: {url}")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Mumbai Public Transport Accessibility | Priority-area and graph-search evaluation")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(115, 115, 115)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = create_paper()
    print(path)
